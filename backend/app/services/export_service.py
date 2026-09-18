"""Eksport raportów do CSV, XLSX i PDF (wszystko po stronie serwera)."""

from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.dates import format_pl_date
from app.core.money import gr_to_zl
from app.models.enums import label as enum_label
from app.services import pdf_service

CSV_DELIMITER = ";"  # zgodne z polskimi ustawieniami Excela
MONEY_FORMAT = "# ##0,00 zł"


def _cell_value(value: object, kind: str) -> object:
    if kind == "money":
        return gr_to_zl(int(value or 0))
    if kind == "date":
        return value if isinstance(value, date) else value
    if kind == "int":
        return int(value or 0)
    if isinstance(value, str):
        return enum_label(value)
    return value


def _csv_value(value: object, kind: str) -> str:
    if kind == "money":
        return f"{gr_to_zl(int(value or 0))}".replace(".", ",")
    if kind == "date" and isinstance(value, date):
        return format_pl_date(value)
    if kind == "int":
        return str(int(value or 0))
    if value is None:
        return ""
    if isinstance(value, str):
        return enum_label(value)
    return str(value)


def report_to_csv(report) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=CSV_DELIMITER, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([title for _, title, _ in report.columns])
    for row in report.rows:
        writer.writerow([_csv_value(row.get(key), kind) for key, _, kind in report.columns])
    if report.summary:
        writer.writerow([])
        for key, value in report.summary.items():
            title = next((col[1] for col in report.columns if col[0] == key), key)
            writer.writerow([f"Razem — {title}", _csv_value(value, "money")])
    # BOM, żeby Excel poprawnie wykrył UTF-8
    return "﻿".encode("utf-8") + buffer.getvalue().encode("utf-8")


def report_to_xlsx(report) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31] or "Raport"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for column_index, (_, title, _) in enumerate(report.columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row_index, row in enumerate(report.rows, start=2):
        for column_index, (key, _, kind) in enumerate(report.columns, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_cell_value(row.get(key), kind))
            if kind == "money":
                cell.number_format = MONEY_FORMAT
            elif kind == "date":
                cell.number_format = "DD.MM.YYYY"

    summary_row = len(report.rows) + 3
    for offset, (key, value) in enumerate(report.summary.items()):
        title = next((col[1] for col in report.columns if col[0] == key), key)
        sheet.cell(row=summary_row + offset, column=1, value=f"Razem — {title}").font = Font(bold=True)
        cell = sheet.cell(row=summary_row + offset, column=2, value=gr_to_zl(int(value)))
        cell.number_format = MONEY_FORMAT
        cell.font = Font(bold=True)

    for column_index, (_, title, kind) in enumerate(report.columns, start=1):
        width = 14 if kind in ("money", "date") else max(12, min(44, len(title) + 8))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_to_pdf(report, *, period_label: str = "") -> bytes:
    return pdf_service.build_report_pdf(report, period_label=period_label)


EXPORT_FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv", report_to_csv),
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
        report_to_xlsx,
    ),
    "pdf": ("application/pdf", "pdf", report_to_pdf),
}
